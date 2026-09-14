# -*- coding: utf-8 -*-
"""交付前自检：申报表数值 / 打印页数 / 类别上限 / 截图完整性 / 条目对应关系。

用法:
  python final_check.py --config check_config.json

check_config.json 示例:
{
  "base_dir": "C:/Users/xx/Desktop/xx综测附加分申报",
  "form": "张三附加分申报表.xlsx",
  "sheet": "Sheet1",
  "header_cells": {"班级": "B3", "姓名": "D3", "学号": "G3"},
  "items": [
    {"row": 10, "tag": "文体活动"},
    {"row": 16, "tag": "各级荣誉-1"},
    {"row": 17, "tag": "各级荣誉-2"},
    {"row": 18, "tag": "学生干部职务"},
    {"row": 24, "tag": "社会实践"}
  ],
  "subtotal_cells": {"学术": "G6", "文体": "G9", "志愿": "G13", "荣誉": "G15",
                     "干部": "G18", "论文": "G20", "其他文章": "G23",
                     "社会实践": "G24", "其他": "G25"},
  "caps": {"学术": 4, "文体": 3, "志愿": 2, "荣誉": 3, "其他文章": 1, "总分": 10},
  "total_cell": "B26",
  "print_area": "A1:G27",
  "shots_dir": "附加分来源截图"
}
"""
import argparse, json, os, tempfile
from openpyxl import load_workbook
from PIL import Image


def check(cfg):
    base = cfg['base_dir']
    rep = []

    # ---- 1) 申报表 ----
    form = os.path.join(base, cfg['form'])
    wb = load_workbook(form, data_only=True)
    ws = wb[cfg.get('sheet', wb.sheetnames[0])]
    rep.append('申报表: %s' % cfg['form'])
    for tag, addr in (cfg.get('header_cells') or {}).items():
        rep.append('  %s = %r' % (tag, ws[addr].value))

    rep.append('  -- 已填条目 --')
    n = 0
    for it in cfg.get('items', []):
        r = it['row']
        vals = [ws.cell(r, c).value for c in range(2, 8)]
        if any(v is not None for v in vals):
            rep.append('  [%s] %s' % (it['tag'], ' | '.join('' if v is None else str(v) for v in vals)))
            n += 1
    rep.append('  已填条目数 = %d' % n)

    subs, caps = cfg.get('subtotal_cells') or {}, cfg.get('caps') or {}
    rep.append('  -- 各类小计 --')
    for tag, addr in subs.items():
        v = ws[addr].value or 0
        cap = caps.get(tag)
        flag = ''
        if cap is not None and float(v) > float(cap):
            flag = '  <<< 超上限 %s！' % cap
        rep.append('  %-8s %s%s' % (tag, v, flag))

    total_addr = cfg.get('total_cell', 'B26')
    total = ws[total_addr].value or 0
    ssum = sum(float(ws[a].value or 0) for a in subs.values())
    rep.append('  总分(%s) = %s   各类相加 = %s   %s'
               % (total_addr, total, round(ssum, 4),
                  'OK' if abs(float(total) - ssum) < 1e-6 else '<<< 不一致！'))
    cap_total = caps.get('总分')
    if cap_total is not None and float(total) > float(cap_total):
        rep.append('  <<< 总分超上限 %s，需按上限截断！' % cap_total)

    # ---- 2) 打印页数 ----
    if cfg.get('verify_pages', True):
        try:
            import tempfile
            import win32com.client as win32
            import pymupdf
            # 临时 PDF 一律放系统临时目录，避免在交付包里建出 .workbuddy 之类的杂质目录
            pdf = os.path.join(tempfile.gettempdir(), 'extra_points_page_check.pdf')
            xl = win32.DispatchEx('Excel.Application')
            xl.DisplayAlerts = False; xl.Visible = False; xl.ScreenUpdating = False
            try:
                # ReadOnly=True：即使用户正开着这个文件也能导出，不会互相锁
                wbx = xl.Workbooks.Open(form, ReadOnly=True)
                wsx = wbx.Worksheets.Item(1)
                wsx.PageSetup.Zoom = False
                wsx.PageSetup.FitToPagesWide = 1
                wsx.PageSetup.FitToPagesTall = 1
                if cfg.get('print_area'):
                    wsx.PageSetup.PrintArea = cfg['print_area']
                wbx.ExportAsFixedFormat(0, pdf)
                wbx.Close(False)
            finally:
                try: xl.Quit()
                finally:
                    try: del wsx, wbx
                    except Exception: pass
                    import gc; gc.collect()
                    try: win32.pythoncom.CoUninitialize()
                    except Exception: pass
                    del xl; gc.collect()
            d = pymupdf.open(pdf)
            rep.append('\n打印页数 = %d %s' % (d.page_count, 'OK' if d.page_count == 1 else '<<< 超过一页！'))
            d.close()
            try: os.unlink(pdf)
            except Exception: pass
        except Exception as e:
            msg = str(e)
            if '可能已被打开' in msg or '无法保存' in msg:
                rep.append('\n打印页数检查失败：文件可能正被 Excel 打开（跳过此项，不影响其它检查）。')
            else:
                rep.append('\n打印页数检查失败: %s' % msg)

    # ---- 3) 截图 ----
    sd = os.path.join(base, cfg.get('shots_dir', '附加分来源截图'))
    files = sorted(os.listdir(sd)) if os.path.isdir(sd) else []
    rep.append('\n截图文件夹「%s」共 %d 个文件' % (os.path.basename(sd), len(files)))
    for f in files:
        fp = os.path.join(sd, f)
        if f.lower().endswith('.png'):
            im = Image.open(fp)
            rep.append('  [PNG] %-46s %s  %.0fKB' % (f, im.size, os.path.getsize(fp) / 1024))
        else:
            rep.append('  [%s] %-44s %.0fKB' % (f.rsplit('.', 1)[-1].upper(), f, os.path.getsize(fp) / 1024))

    out = '\n'.join(rep)
    # 日志默认写系统临时目录；绝不要写进 base_dir，否则会在交付包里多出一个 .workbuddy 目录
    log = cfg.get('log_path') or os.path.join(tempfile.gettempdir(), 'extra_points_final_check.log')
    with open(log, 'w', encoding='utf-8') as fh:
        fh.write(out + '\n')
    print(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--config', required=True)
    a = ap.parse_args()
    check(json.load(open(a.config, encoding='utf-8')))


if __name__ == '__main__':
    main()
