# -*- coding: utf-8 -*-
"""按 JSON 配置填写《附加分申报表》模板：填值 → 公式 → 格式 → 一页打印 → Excel 重算回写。

用法:
  python fill_template.py --config form_config.json [--skip-recalc]

form_config.json 示例:
{
  "template": "C:/.../【附加分】xx学院附加分申报表.xlsx",
  "output":   "C:/.../张三附加分申报表.xlsx",
  "sheet": "Sheet1",
  "header": {"B3": "24xx1班", "D3": "张三", "G3": "202400000001"},
  "header_text_format": {"G3": "@"},
  "rows": [
    {"row": 10, "values": ["北京语言大学校史馆讲解员", "5.5h", "校级", "个人", 0.055]},
    {"row": 16, "values": ["北京语言大学2024-2025学年一二九评优", "先进班集体", "院级", "集体", 0.2]}
  ],
  "value_columns": [2, 3, 4, 5, 6],
  "empty_rows": [6, 7, 8, 11, 12, 14, 20, 21, 23, 25],
  "subtotal_cells": ["G6", "G9", "G13", "G15", "G18", "G20", "G23", "G24", "G25"],
  "total_cell": "B26",
  "total_formula": "=ROUND(G6+G9+G13+G15+G18+G20+G23+G24+G25,3)",
  "number_format": "0.000",
  "print_area": "A1:G27",
  "recalc_with_excel": true,
  "pdf_out": "C:/.../_check.pdf"
}

说明:
- 永远复制模板再改，绝不动原模板。
- 分值列（value_columns 的最后一列）统一用 number_format，避免出现「1.」这种显示。
- 默认用 Excel COM 重算并保存，把公式缓存值写回文件（部分阅读器不计算公式）。
"""
import argparse, json, os, shutil, sys
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side


def fill(cfg):
    src, dst = cfg['template'], cfg['output']
    if os.path.abspath(src) == os.path.abspath(dst):
        raise SystemExit('template 与 output 不能是同一个文件')
    shutil.copyfile(src, dst)

    wb = load_workbook(dst)
    ws = wb[cfg.get('sheet', wb.sheetnames[0])]
    thin = Side(style='thin')
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    nf = cfg.get('number_format', '0.000')
    vcols = cfg.get('value_columns', [2, 3, 4, 5, 6])

    # 表头
    for addr, val in (cfg.get('header') or {}).items():
        ws[addr] = val
    for addr, fmt in (cfg.get('header_text_format') or {}).items():
        ws[addr].number_format = fmt
        ws[addr].font = Font(name='宋体', size=12, bold=True)
        ws[addr].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    def style_row(r, cols):
        for c in cols:
            cell = ws.cell(r, c)
            cell.border = box
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.font = Font(name='宋体', size=11, bold=False)

    # 数据行
    for item in cfg.get('rows', []):
        r, vals = item['row'], item['values']
        for c, v in zip(vcols, vals):
            cell = ws.cell(r, c, v)
            if c == vcols[-1] and isinstance(v, (int, float)):
                cell.number_format = nf
        style_row(r, vcols)

    # 空行（保持边框与居中，视觉上不缺格子）
    for r in cfg.get('empty_rows', []):
        style_row(r, vcols)

    # 小计/总分格式
    for addr in cfg.get('subtotal_cells', []):
        ws[addr].number_format = nf

    # 总分
    if cfg.get('total_formula'):
        ws[cfg['total_cell']] = cfg['total_formula']
        ws[cfg['total_cell']].number_format = nf

    # 打印一页
    pa = cfg.get('print_area')
    if pa:
        ws.print_area = pa
        ws.page_setup.orientation = 'portrait'
        ws.page_setup.paperSize = 9
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.page_setup.scale = None
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.4

    wb.save(dst)
    print('[已生成] %s' % dst)
    return dst


def recalc_and_verify(path, cfg):
    """用 Excel COM 重算并保存，回写公式缓存值；导 PDF 确认页数。"""
    import win32com.client as win32
    xl = win32.DispatchEx('Excel.Application')
    xl.DisplayAlerts = False
    xl.Visible = False
    xl.ScreenUpdating = False
    info = []
    try:
        wb = xl.Workbooks.Open(path)
        ws = wb.Worksheets.Item(1)
        xl.CalculateFullRebuild()
        for addr in list((cfg.get('header') or {}).keys()) + list(cfg.get('subtotal_cells', [])) + \
                [cfg.get('total_cell', 'B26')] + ['%s%d' % (chr(64 + c), it['row'])
                                                  for it in cfg.get('rows', []) for c in cfg.get('value_columns', [])]:
            try:
                info.append('%s = %r' % (addr, ws.Range(addr).Value))
            except Exception:
                pass
        if cfg.get('pdf_out'):
            wb.ExportAsFixedFormat(0, cfg['pdf_out'])
        wb.Save()
        wb.Close(False)
    finally:
        try:
            xl.Quit()
        finally:
            try:
                del ws, wb
            except Exception:
                pass
            import gc
            gc.collect()
            try:
                win32.pythoncom.CoUninitialize()
            except Exception:
                pass
            del xl
            gc.collect()
    print('\n'.join(info))
    return info


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--config', required=True)
    ap.add_argument('--skip-recalc', action='store_true')
    a = ap.parse_args()
    cfg = json.load(open(a.config, encoding='utf-8'))
    out = fill(cfg)
    if cfg.get('recalc_with_excel', True) and not a.skip_recalc:
        recalc_and_verify(out, cfg)
        if cfg.get('pdf_out') and os.path.exists(cfg['pdf_out']):
            import pymupdf
            d = pymupdf.open(cfg['pdf_out'])
            print('\n[页数检查] PDF pages = %d' % d.page_count)
            d.close()


if __name__ == '__main__':
    main()
